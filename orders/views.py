from django.shortcuts import (
    render,
    redirect,
    get_object_or_404
)

from django.contrib.auth.decorators import login_required
from catalog.models import Product
from .cart import Cart
from .models import Order, OrderItem
from .services import generate_order_number
from decimal import Decimal
from django.http import HttpResponse
import openpyxl
from accounts.models import CompanyProfile



@login_required
def add_to_cart(
        request,
        product_id
):

    product = get_object_or_404(
        Product,
        id=product_id
    )

    quantity = Decimal(
        request.POST.get(
            'quantity',
            1
        )
    )

    if quantity < product.min_order_quantity:

        return render(
            request,
            'orders/error.html',
            {
                'message':
                    f'Минимальный заказ: '
                    f'{product.min_order_quantity} '
                    f'{product.unit}'
            }
        )

    if quantity > product.stock:

        return render(
            request,
            'orders/error.html',
            {
                'message':
                    'Недостаточно товара на складе'
            }
        )

    cart = Cart(request)

    cart.add(
        product_id,
        quantity
    )

    return redirect('cart')


@login_required
def cart_view(request):

    cart = Cart(request)

    products = cart.get_products()

    items = []

    for product in products:

        quantity = Decimal(
            str(cart.cart[str(product.id)])
        )

        total = (
            quantity *
            product.price_per_unit
        )

        items.append({
            'product': product,
            'quantity': quantity,
            'total': total
        })

    return render(
        request,
        'orders/cart.html',
        {
            'items': items,
            'total': cart.get_total()
        }
    )


@login_required
def remove_from_cart(
        request,
        product_id
):

    cart = Cart(request)

    cart.remove(product_id)

    return redirect('cart')

@login_required
def create_order(request):

    cart = Cart(request)

    products = cart.get_products()

    if not products.exists():
        return redirect('cart')

    order = Order.objects.create(
        buyer=request.user,
        order_number=generate_order_number()
    )

    total_amount = Decimal('0')

    for product in products:

        quantity = Decimal(
            str(cart.cart[str(product.id)])
        )

        if quantity > product.stock:
            return render(
                request,
                'orders/error.html',
                {
                    'message':
                        f'Недостаточно товара на складе: {product.name}'
                }
            )

        total_price = (
            quantity *
            product.price_per_unit
        )

        OrderItem.objects.create(
            order=order,
            product=product,
            quantity=quantity,
            unit_price=product.price_per_unit,
            total_price=total_price
        )

        product.stock -= quantity
        product.save()

        total_amount += total_price

    order.total_amount = total_amount
    order.save()

    cart.clear()

    return redirect(
        'order_detail',
        order.id
    )

@login_required
def my_orders(request):

    orders = Order.objects.filter(
        buyer=request.user
    ).order_by('-created_at')

    return render(
        request,
        'orders/my_orders.html',
        {
            'orders': orders
        }
    )

@login_required
def order_detail(
        request,
        order_id
):

    order = get_object_or_404(
        Order,
        id=order_id
    )

    return render(
        request,
        'orders/order_detail.html',
        {
            'order': order
        }
    )

@login_required
def export_orders_excel(request):

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "Заказы"

    headers = [
        "Номер заказа",
        "Покупатель",
        "Статус",
        "Сумма",
        "Дата создания"
    ]

    for col_num, header in enumerate(headers, 1):

        sheet.cell(
            row=1,
            column=col_num
        ).value = header

    orders = Order.objects.all()

    row_num = 2

    for order in orders:

        sheet.cell(
            row=row_num,
            column=1
        ).value = order.order_number

        sheet.cell(
            row=row_num,
            column=2
        ).value = order.buyer.username

        sheet.cell(
            row=row_num,
            column=3
        ).value = order.get_status_display()

        sheet.cell(
            row=row_num,
            column=4
        ).value = float(order.total_amount)

        sheet.cell(
            row=row_num,
            column=5
        ).value = order.created_at.strftime(
            "%d.%m.%Y"
        )

        row_num += 1

    response = HttpResponse(
        content_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="orders_report.xlsx"'
    )

    workbook.save(response)

    return response

@login_required
def manage_orders(request):

    profile = request.user.companyprofile

    if profile.role == 'consumer':

        return redirect('/')

    orders = Order.objects.all().order_by(
        '-created_at'
    )

    return render(
        request,
        'orders/manage_orders.html',
        {
            'orders': orders
        }
    )


@login_required
def change_order_status(
        request,
        order_id,
        status
):

    profile = request.user.companyprofile

    if profile.role == 'consumer':

        return redirect('/')

    order = get_object_or_404(
        Order,
        id=order_id
    )

    order.status = status

    order.save()

    return redirect(
        'manage_orders'
    )